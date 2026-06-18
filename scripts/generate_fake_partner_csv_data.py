import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from generate_fake_partner_data import (
    PARTNER_TOTAL,
    extract_template_rows,
    find_excel_sample_row,
    infer_columns,
    partner_record,
    print_stats,
    read_csv_sample,
    trim_trailing_empty_or_zero_columns,
)


def main():
    parser = argparse.ArgumentParser(description="Generate large fake partner CSV data.")
    parser.add_argument("--sample-xlsx", default="input/m4becomvsp_07072024_combine.xlsx")
    parser.add_argument("--sample-csv", default="input/internal_transactions_seed.csv")
    parser.add_argument("--output", default="input/fake_partner_data_realistic.csv")
    args = parser.parse_args()

    sample_xlsx = Path(args.sample_xlsx)
    sample_csv = Path(args.sample_csv)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    _, csv_rows = read_csv_sample(sample_csv)
    wb = load_workbook(sample_xlsx, data_only=True)
    ws = wb.active

    first_data_row = find_excel_sample_row(ws, csv_rows[0])
    columns, base_date = infer_columns(ws, first_data_row, csv_rows[0])
    templates, _ = extract_template_rows(ws, first_data_row)
    required_max_col = max(
        columns["serial"] or 0,
        columns["transaction_id"],
        columns["trace"],
        max(columns["amounts"]),
        max(columns["dates"]),
    )
    templates, max_col = trim_trailing_empty_or_zero_columns(templates, required_max_col)

    with output_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)

        for row_idx in range(1, first_data_row):
            writer.writerow([ws.cell(row_idx, col).value for col in range(1, max_col + 1)])

        for index in range(1, PARTNER_TOTAL + 1):
            record = partner_record(index, base_date)
            row = list(templates[(index - 1) % len(templates)])

            if columns["serial"] is not None:
                row[columns["serial"] - 1] = index
            row[columns["transaction_id"] - 1] = record["transaction_id"]
            row[columns["trace"] - 1] = record["trace"]
            for col in columns["amounts"]:
                row[col - 1] = record["amount"]
            for col in columns["dates"]:
                row[col - 1] = record["trans_date"]

            writer.writerow(row)

    print_stats(output_path)


if __name__ == "__main__":
    main()
