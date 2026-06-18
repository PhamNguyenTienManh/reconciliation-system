import argparse
import csv
import re
from datetime import datetime, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:
    raise SystemExit("Missing dependency: install openpyxl before running this script.") from exc


MATCHED_TOTAL = 99_000
MISMATCHED_TOTAL = 300
PARTNER_ONLY_TOTAL = 300
INTERNAL_ONLY_TOTAL = 400
SHARED_TOTAL = MATCHED_TOTAL + MISMATCHED_TOTAL
PARTNER_TOTAL = SHARED_TOTAL + PARTNER_ONLY_TOTAL
INTERNAL_TOTAL = SHARED_TOTAL + INTERNAL_ONLY_TOTAL

DATE_PATTERN = re.compile(r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}$")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_csv_sample(csv_path):
    with csv_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        header = next(reader)
        rows = [row for row in reader if row]

    if not rows:
        raise ValueError(f"CSV sample has no data rows: {csv_path}")

    return header, rows


def find_excel_sample_row(ws, csv_row):
    csv_values = {value.strip() for value in csv_row if value.strip()}
    best_row = None
    best_score = 0

    for row_idx in range(1, ws.max_row + 1):
        values = {cell_text(ws.cell(row_idx, col).value) for col in range(1, ws.max_column + 1)}
        score = len(csv_values & values)
        if score > best_score:
            best_row = row_idx
            best_score = score

    if best_row is None or best_score < 3:
        raise ValueError("Cannot infer Excel data row from sample Excel and CSV.")

    return best_row


def infer_columns(ws, data_row_idx, csv_row):
    row_values = {
        col: cell_text(ws.cell(data_row_idx, col).value)
        for col in range(1, ws.max_column + 1)
    }

    date_value = next((value for value in csv_row if DATE_PATTERN.match(value.strip())), None)
    if not date_value:
        raise ValueError("Cannot infer date format from CSV sample.")

    amount_value = None
    for value in csv_row:
        normalized = value.strip().replace(",", "")
        if normalized.isdigit() and list(row_values.values()).count(normalized) >= 2:
            amount_value = normalized
            break
    if not amount_value:
        raise ValueError("Cannot infer amount column from sample Excel and CSV.")

    matched_numeric_values = [
        value.strip()
        for value in csv_row
        if value.strip().isdigit()
        and value.strip() in row_values.values()
        and value.strip() != amount_value
    ]
    if len(matched_numeric_values) < 2:
        raise ValueError("Cannot infer transaction id and trace columns from samples.")

    matched_numeric_values.sort(key=len)
    transaction_id_value = matched_numeric_values[0]
    trace_value = matched_numeric_values[-1]

    columns = {
        "serial": None,
        "transaction_id": None,
        "trace": None,
        "amounts": [],
        "dates": [],
    }

    for col, value in row_values.items():
        if columns["serial"] is None and value == "1":
            columns["serial"] = col
        if value == transaction_id_value:
            columns["transaction_id"] = col
        if value == trace_value:
            columns["trace"] = col
        if value == amount_value:
            columns["amounts"].append(col)
        if value == date_value:
            columns["dates"].append(col)

    missing = [key for key in ("transaction_id", "trace") if columns[key] is None]
    if missing or not columns["amounts"] or not columns["dates"]:
        raise ValueError(f"Cannot infer required Excel columns. Missing: {missing}")

    return columns, datetime.strptime(date_value, "%d/%m/%Y %H:%M:%S")


def extract_template_rows(ws, first_data_row):
    rows = []
    max_col = 0
    for row_idx in range(first_data_row, ws.max_row + 1):
        if any(cell_text(ws.cell(row_idx, col).value) for col in range(1, ws.max_column + 1)):
            rows.append(row_idx)
            for col in range(1, ws.max_column + 1):
                if cell_text(ws.cell(row_idx, col).value):
                    max_col = max(max_col, col)

    if not rows:
        raise ValueError("Excel sample has no data rows.")

    templates = []
    for row_idx in rows:
        templates.append([ws.cell(row_idx, col).value for col in range(1, max_col + 1)])

    return templates, max_col


def trim_trailing_empty_or_zero_columns(templates, required_max_col):
    max_col = max(len(row) for row in templates)
    for col_index in range(max_col - 1, required_max_col - 1, -1):
        values = [cell_text(row[col_index]) for row in templates if col_index < len(row)]
        if any(value and value != "0" for value in values):
            break
    else:
        col_index = required_max_col - 1

    trimmed_max_col = col_index + 1
    return [row[:trimmed_max_col] for row in templates], trimmed_max_col


def fake_trace(index, prefix):
    return f"{prefix}{index:019d}"


def fake_transaction_id(index, prefix):
    return f"{prefix}{index:08d}"


def fake_amount(index):
    return 10_000 + (index % 9_000) * 100


def fake_date(base_date, index):
    return (base_date + timedelta(seconds=index)).strftime("%d/%m/%Y %H:%M:%S")


def partner_record(index, base_date):
    if index <= SHARED_TOTAL:
        sequence = index
        trace = fake_trace(sequence, "90")
    else:
        sequence = index - SHARED_TOTAL
        trace = fake_trace(sequence, "91")

    return {
        "transaction_id": fake_transaction_id(index, "70"),
        "trace": trace,
        "amount": fake_amount(sequence),
        "trans_date": fake_date(base_date, sequence),
    }


def print_stats(output_path):
    print(f"Generated partner file: {output_path}")
    print(f"Partner records: {PARTNER_TOTAL:,}")
    print(f"Internal records: {INTERNAL_TOTAL:,}")
    print("Expected reconciliation cases:")
    print(f"  MATCHED: {MATCHED_TOTAL:,}")
    print(f"  MISMATCHED: {MISMATCHED_TOTAL:,}")
    print(f"  PARTNER_ONLY: {PARTNER_ONLY_TOTAL:,}")
    print(f"  INTERNAL_ONLY: {INTERNAL_ONLY_TOTAL:,}")


def main():
    parser = argparse.ArgumentParser(description="Generate large fake partner Excel data.")
    parser.add_argument("--sample-xlsx", default="input/m4becomvsp_07072024_combine.xlsx")
    parser.add_argument("--sample-csv", default="input/internal_transactions_seed.csv")
    parser.add_argument("--output", default="input/fake_partner_data_realistic.xlsx")
    args = parser.parse_args()

    sample_xlsx = Path(args.sample_xlsx)
    sample_csv = Path(args.sample_csv)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    _, csv_rows = read_csv_sample(sample_csv)
    wb = load_workbook(sample_xlsx)
    ws = wb.active

    first_data_row = find_excel_sample_row(ws, csv_rows[0])
    columns, base_date = infer_columns(ws, first_data_row, csv_rows[0])
    templates, max_col = extract_template_rows(ws, first_data_row)
    required_max_col = max(
        columns["serial"] or 0,
        columns["transaction_id"],
        columns["trace"],
        max(columns["amounts"]),
        max(columns["dates"]),
    )
    templates, max_col = trim_trailing_empty_or_zero_columns(templates, required_max_col)

    output_wb = Workbook(write_only=True)
    output_ws = output_wb.create_sheet(title=ws.title)

    for row_idx in range(1, first_data_row):
        output_ws.append([ws.cell(row_idx, col).value for col in range(1, max_col + 1)])

    for index in range(1, PARTNER_TOTAL + 1):
        record = partner_record(index, base_date)
        row = list(templates[(index - 1) % len(templates)])

        if columns["serial"] is not None:
            row[columns["serial"] - 1] = index
        row[columns["transaction_id"] - 1] = record["transaction_id"]
        row[columns["trace"] - 1] = record["trace"]
        for col in columns["amounts"]:
            row[col - 1] = record["amount"]
        for col in columns["dates"]:
            row[col - 1] = record["trans_date"]
        output_ws.append(row)

    output_wb.save(output_path)
    print_stats(output_path)


if __name__ == "__main__":
    main()
